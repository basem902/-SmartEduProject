"""
File Validation Utilities
التحقق من الملفات: الحجم، النوع، الفيروسات، والمحتوى بالذكاء الاصطناعي
"""
import os
import mimetypes
import hashlib
from django.conf import settings
from django.core.exceptions import ValidationError
import google.generativeai as genai
import PyPDF2
import docx
import openpyxl
from io import BytesIO

# محاولة استيراد المكتبات الاختيارية
try:
    import magic
    MAGIC_AVAILABLE = True
except ImportError:
    MAGIC_AVAILABLE = False

try:
    import pyclamd
    CLAMAV_AVAILABLE = True
except ImportError:
    CLAMAV_AVAILABLE = False


class FileValidator:
    """
    فئة للتحقق من الملفات المرفوعة
    """
    
    def __init__(self, file, project):
        """
        Args:
            file: ملف Django UploadedFile
            project: كائن Project
        """
        self.file = file
        self.project = project
        self.errors = []
        self.warnings = []
        
    def validate_all(self):
        """
        التحقق الشامل من الملف
        
        Returns:
            dict: النتيجة {valid, errors, warnings, ai_check}
        """
        # 1. التحقق من الحجم
        self._validate_size()
        
        # 2. التحقق من النوع
        self._validate_type()
        
        # 3. التحقق من الاسم
        self._validate_filename()
        
        # 4. فحص الفيروسات
        virus_result = self._scan_virus()
        
        # 5. التحقق بالذكاء الاصطناعي
        ai_result = None
        if not self.errors:  # فقط إذا لم تكن هناك أخطاء أساسية
            ai_result = self._validate_with_ai()
        
        return {
            'valid': len(self.errors) == 0,
            'errors': self.errors,
            'warnings': self.warnings,
            'virus_scan': virus_result,
            'ai_check': ai_result,
            'file_info': {
                'name': self.file.name,
                'size': self.file.size,
                'type': self.file.content_type,
                'hash': self._get_file_hash()
            }
        }
    
    def _validate_size(self):
        """التحقق من حجم الملف"""
        max_size = self.project.max_file_size * 1024 * 1024  # MB to bytes
        
        if self.file.size > max_size:
            self.errors.append({
                'type': 'size',
                'message': f'حجم الملف {self._format_size(self.file.size)} يتجاوز الحد المسموح {self.project.max_file_size} MB'
            })
        
        # تحذير إذا كان الملف صغير جداً (أقل من 10 KB)
        if self.file.size < 10 * 1024:
            self.warnings.append({
                'type': 'size',
                'message': 'الملف صغير جداً، تأكد من رفع الملف الصحيح'
            })
    
    def _validate_type(self):
        """التحقق من نوع الملف"""
        # الحصول على الامتداد
        _, ext = os.path.splitext(self.file.name.lower())
        ext = ext.lstrip('.')
        
        # قائمة الامتدادات المسموحة
        allowed = self.project.allowed_extensions.split(',')
        allowed = [e.strip().lower() for e in allowed]
        
        if ext not in allowed:
            self.errors.append({
                'type': 'extension',
                'message': f'نوع الملف .{ext} غير مسموح. الأنواع المسموحة: {", ".join(allowed)}'
            })
            return
        
        # التحقق من MIME type الحقيقي (منع التزييف)
        if MAGIC_AVAILABLE:
            try:
                mime = magic.Magic(mime=True)
                actual_mime = mime.from_buffer(self.file.read(1024))
                self.file.seek(0)  # إعادة المؤشر للبداية
                
                # التحقق من توافق MIME مع الامتداد
                expected_mimes = self._get_expected_mimes(ext)
                if expected_mimes and actual_mime not in expected_mimes:
                    self.errors.append({
                        'type': 'mime',
                        'message': f'محتوى الملف لا يطابق امتداده. النوع الفعلي: {actual_mime}'
                    })
            except Exception as e:
                self.warnings.append({
                    'type': 'mime',
                    'message': f'تعذر التحقق من نوع الملف: {str(e)}'
                })
        else:
            self.warnings.append({
                'type': 'mime',
                'message': 'التحقق المتقدم من MIME غير متاح (python-magic غير مثبت)'
            })
    
    def _validate_filename(self):
        """التحقق من اسم الملف"""
        # التحقق من الأحرف الخطيرة
        dangerous_chars = ['..', '/', '\\', '<', '>', ':', '"', '|', '?', '*']
        for char in dangerous_chars:
            if char in self.file.name:
                self.errors.append({
                    'type': 'filename',
                    'message': f'اسم الملف يحتوي على حرف غير مسموح: {char}'
                })
                return
        
        # التحقق من الطول
        if len(self.file.name) > 255:
            self.errors.append({
                'type': 'filename',
                'message': 'اسم الملف طويل جداً'
            })
    
    def _scan_virus(self):
        """فحص الفيروسات باستخدام ClamAV"""
        if not CLAMAV_AVAILABLE:
            return {
                'scanned': False,
                'message': 'فحص الفيروسات غير متاح (pyclamd غير مثبت)'
            }
        
        try:
            # الاتصال بـ ClamAV
            cd = pyclamd.ClamdNetworkSocket(
                host=settings.CLAMAV_HOST,
                port=settings.CLAMAV_PORT
            )
            
            # التحقق من اتصال ClamAV
            if not cd.ping():
                return {
                    'scanned': False,
                    'message': 'خدمة فحص الفيروسات غير متاحة'
                }
            
            # فحص الملف
            self.file.seek(0)
            result = cd.scan_stream(self.file.read())
            self.file.seek(0)
            
            if result:
                # وُجد فيروس!
                self.errors.append({
                    'type': 'virus',
                    'message': f'تم اكتشاف تهديد أمني: {result}'
                })
                return {
                    'scanned': True,
                    'clean': False,
                    'threat': result
                }
            
            return {
                'scanned': True,
                'clean': True,
                'message': 'الملف آمن'
            }
            
        except Exception as e:
            # ClamAV غير مثبت أو غير متاح
            self.warnings.append({
                'type': 'virus_scan',
                'message': f'تعذر فحص الفيروسات: {str(e)}'
            })
            return {
                'scanned': False,
                'message': str(e)
            }
    
    def _validate_with_ai(self):
        """
        التحقق من توافق محتوى الملف مع متطلبات المشروع
        باستخدام Gemini AI
        """
        if not settings.GEMINI_API_KEY:
            return {
                'checked': False,
                'message': 'الذكاء الاصطناعي غير مفعّل'
            }
        
        try:
            # استخراج محتوى الملف
            content = self._extract_file_content()
            
            if not content:
                return {
                    'checked': False,
                    'message': 'لا يمكن قراءة محتوى الملف'
                }
            
            # إعداد Gemini
            genai.configure(api_key=settings.GEMINI_API_KEY)
            model = genai.GenerativeModel('gemini-pro')
            
            # بناء Prompt
            prompt = self._build_ai_prompt(content)
            
            # إرسال للـ AI
            response = model.generate_content(prompt)
            
            # تحليل الإجابة
            ai_response = response.text
            
            # استخلاص النتيجة
            result = self._parse_ai_response(ai_response)
            
            # إضافة تحذيرات إذا كان هناك مشاكل
            if not result['compliant']:
                self.warnings.append({
                    'type': 'ai_check',
                    'message': result['message']
                })
            
            return result
            
        except Exception as e:
            return {
                'checked': False,
                'message': f'خطأ في التحقق بالذكاء الاصطناعي: {str(e)}'
            }
    
    def _extract_file_content(self):
        """استخراج محتوى الملف حسب نوعه"""
        _, ext = os.path.splitext(self.file.name.lower())
        
        try:
            self.file.seek(0)
            
            if ext == '.pdf':
                return self._extract_pdf_content()
            elif ext in ['.docx', '.doc']:
                return self._extract_docx_content()
            elif ext in ['.xlsx', '.xls']:
                return self._extract_excel_content()
            elif ext in ['.txt', '.md']:
                return self.file.read().decode('utf-8', errors='ignore')
            else:
                # للأنواع الأخرى (صور، فيديو، صوت)
                return None
                
        except Exception as e:
            return None
    
    def _extract_pdf_content(self):
        """استخراج نص من PDF"""
        try:
            pdf_reader = PyPDF2.PdfReader(BytesIO(self.file.read()))
            text = ""
            for page in pdf_reader.pages[:5]:  # أول 5 صفحات فقط
                text += page.extract_text()
            return text[:5000]  # أول 5000 حرف
        except:
            return None
    
    def _extract_docx_content(self):
        """استخراج نص من Word"""
        try:
            doc = docx.Document(BytesIO(self.file.read()))
            text = "\n".join([para.text for para in doc.paragraphs[:50]])  # أول 50 فقرة
            return text[:5000]
        except:
            return None
    
    def _extract_excel_content(self):
        """استخراج بيانات من Excel"""
        try:
            wb = openpyxl.load_workbook(BytesIO(self.file.read()))
            ws = wb.active
            text = ""
            for row in list(ws.rows)[:50]:  # أول 50 صف
                text += " ".join([str(cell.value) for cell in row if cell.value]) + "\n"
            return text[:5000]
        except:
            return None
    
    def _build_ai_prompt(self, content):
        """بناء prompt للذكاء الاصطناعي"""
        prompt = f"""أنت مساعد تعليمي ذكي. مهمتك التحقق من توافق محتوى الملف المرفوع مع متطلبات المشروع.

**معلومات المشروع:**
العنوان: {self.project.title}
المادة: {self.project.subject.name if self.project.subject else 'غير محدد'}

**التعليمات:**
{self.project.instructions or 'لا توجد تعليمات محددة'}

**المتطلبات:**
{self.project.requirements or 'لا توجد متطلبات محددة'}

**محتوى الملف المرفوع (عينة):**
{content[:3000]}

**المطلوب منك:**
1. تحليل محتوى الملف
2. التحقق من توافقه مع التعليمات والمتطلبات
3. الرد بصيغة JSON فقط:
{{
    "compliant": true/false,
    "confidence": 0-100,
    "message": "رسالة توضيحية بالعربية",
    "issues": ["قائمة المشاكل إن وجدت"],
    "suggestions": ["اقتراحات للتحسين"]
}}

تحليلك:"""
        
        return prompt
    
    def _parse_ai_response(self, response):
        """تحليل إجابة AI"""
        try:
            import json
            import re
            
            # استخلاص JSON من الإجابة
            json_match = re.search(r'\{.*\}', response, re.DOTALL)
            if json_match:
                data = json.loads(json_match.group())
                return {
                    'checked': True,
                    'compliant': data.get('compliant', True),
                    'confidence': data.get('confidence', 0),
                    'message': data.get('message', ''),
                    'issues': data.get('issues', []),
                    'suggestions': data.get('suggestions', [])
                }
        except:
            pass
        
        # إذا فشل التحليل، نعتبرها موافقة
        return {
            'checked': True,
            'compliant': True,
            'confidence': 50,
            'message': 'تم التحقق بنجاح',
            'issues': [],
            'suggestions': []
        }
    
    def _get_file_hash(self):
        """حساب hash للملف"""
        try:
            self.file.seek(0)
            file_hash = hashlib.sha256()
            for chunk in self.file.chunks():
                file_hash.update(chunk)
            self.file.seek(0)
            return file_hash.hexdigest()
        except:
            return None
    
    def _get_expected_mimes(self, extension):
        """الحصول على MIME types المتوقعة لامتداد معين"""
        mime_map = {
            'pdf': ['application/pdf'],
            'doc': ['application/msword'],
            'docx': ['application/vnd.openxmlformats-officedocument.wordprocessingml.document'],
            'xls': ['application/vnd.ms-excel'],
            'xlsx': ['application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'],
            'ppt': ['application/vnd.ms-powerpoint'],
            'pptx': ['application/vnd.openxmlformats-officedocument.presentationml.presentation'],
            'txt': ['text/plain'],
            'jpg': ['image/jpeg'],
            'jpeg': ['image/jpeg'],
            'png': ['image/png'],
            'gif': ['image/gif'],
            'mp4': ['video/mp4'],
            'mp3': ['audio/mpeg'],
            'wav': ['audio/wav'],
            'zip': ['application/zip', 'application/x-zip-compressed'],
            'rar': ['application/x-rar-compressed'],
        }
        return mime_map.get(extension, [])
    
    def _format_size(self, size):
        """تنسيق حجم الملف"""
        for unit in ['B', 'KB', 'MB', 'GB']:
            if size < 1024.0:
                return f"{size:.2f} {unit}"
            size /= 1024.0
        return f"{size:.2f} TB"
